from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path("/Users/wangshihan/Air Quality Research")
OUT_DIR = ROOT / "Reports"
OUT_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "FS12联邦学习论文课题与实验进度报告.docx"
FIGURE_PATH = Path("/var/folders/sb/bws0mwjj1vddxfsv5vsdlj500000gn/T/codex-clipboard-4efa13fd-9688-4199-abb7-c392b57b8f95.png")

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F5F9FC"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
TEXT = RGBColor(32, 32, 32)
MUTED = RGBColor(90, 90, 90)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = tc_borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, font="Arial Unicode MS"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_para_font(paragraph, size=11, color=TEXT, bold=None):
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bold)


def add_body(doc, text, bold_lead=None, after=6, first_line=True):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.35
    if first_line:
        p.paragraph_format.first_line_indent = Pt(22)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=11, bold=True, color=TEXT)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=11, color=TEXT)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=TEXT)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(item)
        set_run_font(r, size=10.8, color=TEXT)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Pt(22)
        p.paragraph_format.first_line_indent = Pt(-22)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(f"{index}.  {item}")
        set_run_font(r, size=10.8, color=TEXT)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    for r in p.runs:
        set_run_font(r, size=16 if level == 1 else 12.5, bold=True, color=RGBColor(0, 0, 0))
    return p


def add_table(doc, headers, rows, widths, numeric_cols=None, font_size=9.5):
    numeric_cols = set(numeric_cols or [])
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, (cell, label, width) in enumerate(zip(hdr.cells, headers, widths)):
        set_col_width(cell, width)
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell, top=130, bottom=130, start=110, end=110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_run_font(r, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, (cell, value, width) in enumerate(zip(cells, row, widths)):
            set_col_width(cell, width)
            set_cell_borders(cell)
            set_cell_margins(cell, top=115, bottom=115, start=110, end=110)
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cidx in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, color=MUTED)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial Unicode MS"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
normal.font.size = Pt(11)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(6)

for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
    style = styles[style_name]
    style.font.name = "Arial Unicode MS"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    style.font.color.rgb = RGBColor(0, 0, 0)

# Remove the decorative border inherited by some Word Title themes.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

footer = section.footer.paragraphs[0]
add_page_number(footer)

# Cover page
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(72)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("FS12联邦学习论文课题与实验进度报告")
set_run_font(r, size=25, bold=True, color=RGBColor(0, 0, 0))
p.style = styles["Title"]

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(28)
r = p.add_run("面向教授讨论的阶段性分析")
set_run_font(r, size=14, color=MUTED)

meta = add_table(
    doc,
    ["项目", "内容"],
    [
        ["研究任务", "13个监测站的小时级 PM2.5 联邦预测"],
        ["预测设置", "回看 168 小时并预测未来 24 小时"],
        ["比较对象", "Transformer 与 Mamba 结合 FedAvg 与 FedProx"],
        ["当前实验规模", "4组实验 每组30轮通信"],
        ["报告日期", "2026年9月20日"],
    ],
    [1.45, 5.45],
    font_size=10.5,
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(26)
p.paragraph_format.left_indent = Inches(0.25)
p.paragraph_format.right_indent = Inches(0.25)
p.paragraph_format.line_spacing = 1.35
r = p.add_run(
    "当前结果显示，Mamba 与 FedAvg 的组合在预测精度上略优，同时将单轮训练时间缩短到 Transformer 的约五分之一。"
    "这一结果为论文形成了清晰的效率主线，但精度差异仍然很小，现阶段应把结论表述为初步证据，并优先补充多随机种子、验证集、统计检验和公平性审计。"
)
set_run_font(r, size=12, color=TEXT)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(36)
r = p.add_run("供课题讨论与下一阶段实验决策使用")
set_run_font(r, size=10, color=MUTED)

doc.add_page_break()

add_heading(doc, "执行摘要", 1)
add_body(
    doc,
    "当前四组实验已经完成了第一轮可比较的结果。就现有数据而言，Mamba 加 FedAvg 是最值得继续推进的主方案。"
    "其最终13站平均 MAE 为 1.9021，RMSE 为 4.3301，R² 为 0.2744；收敛曲线中的最低平均测试 MAE 为 1.8992。"
    "与 Transformer 加 FedAvg 相比，精度改善约为 0.15%，但每轮训练速度约快 4.82 倍，总实验时间约缩短 79%。",
)
add_body(
    doc,
    "目前最有说服力的是计算效率，而不是精度优势。四种方法的 MAE 差距不超过 0.003，单次随机种子不足以证明差异稳定。"
    "FedProx 在当前设置下没有带来一致收益，说明简单的 FedAvg 已能处理当前程度的站点差异，或者 FedProx 的系数和本地训练轮数尚未进入适合它发挥作用的区间。",
)
add_body(
    doc,
    "下一阶段不建议立即扩大模型数量。应先完成研究可信度所必需的实验：使用独立验证集进行早停和模型选择；至少运行三个随机种子；报告站点级差异和置信区间；核对参数量、训练深度、混合精度和硬件条件。"
    "完成这些工作后，再开展序列长度、预测范围和数据异质性实验，检验 Mamba 的效率优势是否随任务难度稳定存在。",
)

add_heading(doc, "建议向教授汇报的核心观点", 2)
add_bullets(doc, [
    "已经获得一个清晰的初步结果：Mamba 加 FedAvg 在当前任务中同时实现了略好的精度和显著更低的训练时间。",
    "现阶段可以主张明显的效率优势，但不能主张显著的精度提升，因为差值很小且尚未进行重复实验和统计检验。",
    "FedProx 没有明显优于 FedAvg。该结果本身有研究价值，但需要通过异质性指标和超参数实验解释原因。",
    "论文的下一步重点应从继续堆模型转向验证结果是否稳定、公平且可复现。",
])

doc.add_page_break()
add_heading(doc, "一 研究课题与当前实验设计", 1)
add_body(
    doc,
    "本课题研究多个空气质量监测站在不共享原始数据的条件下，如何共同训练 PM2.5 预测模型。每个监测站作为一个联邦客户端，模型在本地数据上更新参数，服务器只聚合模型参数。"
    "当前实验同时比较两类时序骨干网络和两类联邦聚合方法，以回答模型结构和聚合策略分别带来多大影响。",
)

add_table(
    doc,
    ["维度", "当前设置", "研究含义"],
    [
        ["数据", "13个小时级监测站", "模拟跨站点且数据分布不同的联邦场景"],
        ["输入", "过去168小时", "使用最近7天信息预测未来变化"],
        ["输出", "未来24小时 PM2.5", "对应一天的短期空气质量预测"],
        ["特征", "FS12协变量与 PM2.5", "同时使用污染物和气象信息"],
        ["骨干网络", "Transformer 与 Mamba", "比较注意力模型和状态空间模型"],
        ["聚合方法", "FedAvg 与 FedProx", "比较标准参数平均和近端约束"],
        ["训练", "30轮通信", "每轮各客户端本地训练后聚合"],
        ["指标", "MAE RMSE R² 耗时", "同时评价准确性与计算成本"],
    ],
    [1.25, 2.15, 3.45],
    font_size=9.3,
)

add_heading(doc, "当前论文问题", 2)
add_body(
    doc,
    "当前结果最适合支持的研究问题是：在多站点 PM2.5 联邦预测中，Mamba 能否在保持预测精度的同时降低训练时间；FedProx 是否能在站点数据分布不同的情况下进一步改善模型。"
    "如果后续实验确认这一趋势，论文可以围绕准确性与效率的权衡展开，而不是只比较哪个模型的单一误差最低。",
)

doc.add_page_break()
add_heading(doc, "二 当前实验结果", 1)

add_heading(doc, "最终13站平均指标", 2)
add_table(
    doc,
    ["实验", "MAE", "RMSE", "R²", "每轮耗时", "总时长"],
    [
        ["Transformer 加 FedAvg", "1.9049", "4.3368", "0.2719", "167.8秒", "1.67小时"],
        ["Transformer 加 FedProx", "1.9039", "4.3363", "0.2720", "178.3秒", "1.76小时"],
        ["Mamba 加 FedAvg", "1.9021", "4.3301", "0.2744", "34.8秒", "0.35小时"],
        ["Mamba 加 FedProx", "1.9029", "4.3392", "0.2702", "38.5秒", "0.39小时"],
    ],
    [2.05, 0.72, 0.82, 0.72, 1.02, 0.94],
    numeric_cols={1, 2, 3, 4, 5},
    font_size=8.7,
)
add_caption(doc, "表1  四种方法在第30轮后的13站平均结果")

add_heading(doc, "收敛和速度指标", 2)
add_table(
    doc,
    ["实验", "进入最优值1%区间", "曲线最低 MAE", "相对同聚合 Transformer 的单轮速度"],
    [
        ["Transformer 加 FedAvg", "第15轮", "1.9049", "基准"],
        ["Transformer 加 FedProx", "第15轮", "1.9039", "基准"],
        ["Mamba 加 FedAvg", "第10轮", "1.8992", "约4.82倍"],
        ["Mamba 加 FedProx", "第15轮", "1.9024", "约4.63倍"],
    ],
    [2.05, 1.45, 1.18, 2.00],
    numeric_cols={1, 2, 3},
    font_size=9.0,
)
add_caption(doc, "表2  收敛轮数表示首次进入全程最低 MAE 的1%范围 并不等同于损失已经稳定")

if FIGURE_PATH.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(FIGURE_PATH), width=Inches(6.75))
    add_caption(doc, "图1  四种方法在第5至30轮的13站平均测试 MAE")

add_heading(doc, "三 结果分析", 1)

add_heading(doc, "Mamba 加 FedAvg 提供最佳综合表现", 2)
add_body(
    doc,
    "Mamba 加 FedAvg 在最终13站平均结果中取得最低 MAE 和 RMSE，同时得到最高 R²。"
    "相对 Transformer 加 FedAvg，其 MAE 下降 0.0028，约为 0.15%；RMSE 下降 0.0067，约为 0.15%；R² 提高 0.0025。"
    "这些变化方向一致，但幅度很小，因此目前应称为轻微改善，而不是显著提升。",
)

add_heading(doc, "训练效率是当前最清晰的结果", 2)
add_body(
    doc,
    "在相同聚合方法下，Mamba 的单轮训练速度分别约为 Transformer 的 4.82 倍和 4.63 倍。"
    "总时长分别从 1.67 小时降到 0.35 小时，以及从 1.76 小时降到 0.39 小时，缩短约 77% 至 79%。"
    "按照平均单轮耗时粗略估算，Mamba 加 FedAvg 进入最优值1%范围约需 5.8 分钟，而 Transformer 加 FedAvg 约需 42.0 分钟。"
    "这个时间估计未包含每次评估和初始化成本，但足以说明两者在实验吞吐量上的明显差别。",
)

add_heading(doc, "FedProx 暂未表现出稳定优势", 2)
add_body(
    doc,
    "Transformer 使用 FedProx 后，MAE 只改善 0.0010，其他指标几乎不变，但每轮耗时增加约 6.3%。"
    "Mamba 使用 FedProx 后，最终 MAE 比 FedAvg 高 0.0008，RMSE 高 0.0091，R² 低 0.0042，而且训练更慢。"
    "这说明在当前 μ 等于 0.01、本地训练1轮的条件下，近端约束没有提供足以抵消额外成本的收益。",
)

add_heading(doc, "Mamba 在后期出现轻微回升", 2)
add_body(
    doc,
    "Mamba 加 FedAvg 在第25轮达到最低测试 MAE 1.8992，第30轮回升到约 1.9021。"
    "回升幅度约为 0.15%，可能来自轻微过拟合、客户端更新波动或随机噪声。"
    "这一现象提示后续应使用验证集保存最佳轮次，而不应固定使用最后一轮，也不应根据测试集选择模型。",
)

add_heading(doc, "当前 R² 表明仍有较大提升空间", 2)
add_body(
    doc,
    "四种模型的平均 R² 约为 0.27，说明模型能够解释部分 PM2.5 波动，但大量短时变化仍未被现有输入捕捉。"
    "可能原因包括突发污染事件、火点和烟霾传输、交通与工业排放、风向的空间效应以及缺失值处理。"
    "在加入朴素基线和集中式上限之前，不能仅凭 R² 判断当前模型是否已经达到有实际意义的预测水平。",
)

doc.add_page_break()
add_heading(doc, "四 当前结论的可信范围", 1)

add_heading(doc, "可以明确表达的结论", 2)
add_bullets(doc, [
    "在本次 RTX 4090 实验、当前数据划分和当前超参数下，Mamba 的训练时间显著短于 Transformer。",
    "四种方法的最终平均精度非常接近，Mamba 加 FedAvg 暂时排名第一。",
    "FedProx 在当前设置下没有稳定改善准确率，并增加了训练时间。",
    "Mamba 加 FedAvg 在第25轮取得最低曲线值，继续训练到第30轮没有进一步改善。",
])

add_heading(doc, "暂时不能表达的结论", 2)
add_bullets(doc, [
    "不能说 Mamba 的预测精度显著优于 Transformer，因为目前看起来只有一个随机种子，差距也小于 0.2%。",
    "不能说 FedProx 对空气质量联邦学习无效。当前只测试了一个 μ 和一个本地训练轮数。",
    "不能把第10轮或第15轮直接称为模型已经收敛。当前指标只是首次进入全程最佳值的1%范围。",
    "不能只凭13站平均值说明所有站点都得到改善。平均值可能掩盖少数站点的明显退化。",
    "不能在缺少参数量和实现审计时，把全部速度差异归因于 Mamba 架构本身。",
])

add_heading(doc, "需要优先处理的方法学问题", 2)
add_table(
    doc,
    ["问题", "对论文的影响", "建议处理"],
    [
        ["测试集被用于每5轮绘制曲线和寻找最低值", "如果据此早停或选择模型 会形成测试信息泄漏", "按时间建立独立验证集 只在最终锁定模型后评估测试集"],
        ["只有单次随机种子", "无法区分模型差异和初始化波动", "至少运行2023 2024 2025三个种子并报告均值和标准差"],
        ["收敛阈值为最低 MAE 的1%", "阈值约0.019 大于模型间差距 因而偏宽松", "改名为进入近最优区间 或使用验证集耐心值与更严格阈值"],
        ["只报告13站平均值", "无法判断改进是否普遍存在", "报告每站结果 中位数 四分位数 胜出站点数和最差站点"],
        ["模型复杂度尚未对齐", "速度优势可能同时受到层数 参数量和内核实现影响", "报告参数量 FLOPs 峰值显存 混合精度和实际执行代码版本"],
        ["缺少朴素和训练方式基线", "无法判断联邦学习是否真正增加价值", "加入持续性 季节性朴素 本地独立和集中式训练"],
    ],
    [1.9, 2.35, 2.65],
    font_size=8.7,
)

add_heading(doc, "五 下一阶段实验计划", 1)
add_body(
    doc,
    "下一阶段应先提高结论可信度，再扩展研究问题。下面的顺序按对论文有效性的影响排序。第一阶段完成后，当前结果才适合用于论文主表；第二阶段用于解释机制和形成论文贡献；第三阶段用于完善工程与成本分析。",
)

add_table(
    doc,
    ["优先级", "工作", "具体做法", "预期输出"],
    [
        ["第一", "建立验证方案", "按时间划分训练 验证 测试并设置无重叠间隔 测试集只评估一次", "可信的最佳轮次和最终测试结果"],
        ["第一", "完成多种子复现", "四种方法至少运行3个相同种子 保持数据划分一致", "均值 标准差 95%置信区间"],
        ["第一", "进行站点级统计", "对13站做配对比较 报告胜出站点数并使用配对检验", "差异是否稳定以及哪些站点受益"],
        ["第一", "核对公平性与复现性", "冻结实际运行代码 环境和配置 记录参数量 显存 GPU利用率", "可审计的实验配置表"],
        ["第二", "加入必要基线", "持续性 季节性朴素 本地独立 集中式训练", "证明预测模型和联邦训练的增量价值"],
        ["第二", "分析 FedProx", "测试 μ 为0.001 0.01 0.1以及本地轮数1 3 5", "解释 FedProx 在何种异质性下有效"],
        ["第二", "序列和预测范围消融", "回看24 72 168 336小时 预测6 12 24 48小时", "确认 Mamba 优势是否随序列变长扩大"],
        ["第二", "特征消融", "比较仅 PM2.5 气象特征和完整 FS12", "量化协变量对误差和 R² 的贡献"],
        ["第三", "效率和成本分析", "记录吞吐量 峰值显存 功耗或租用成本", "精度 时间 显存和成本的综合比较"],
    ],
    [0.65, 1.25, 3.05, 1.95],
    font_size=8.3,
)

add_heading(doc, "建议的统计报告方式", 2)
add_bullets(doc, [
    "主表报告三个或更多随机种子的站点平均结果，格式为均值加减标准差。",
    "对同一站点和同一种子形成配对结果，使用配对 t 检验或 Wilcoxon 符号秩检验，并报告效应量与95%置信区间。",
    "同时报告宏平均和按测试样本数加权的微平均，避免站点规模不同造成解释偏差。",
    "使用验证集决定最佳轮次，最终测试集只产生一次 MAE RMSE 和 R²。",
])

add_heading(doc, "六 论文定位与可检验观点", 1)

add_heading(doc, "推荐的论文主线", 2)
add_body(
    doc,
    "现阶段最稳妥的论文定位是高效的多站点联邦 PM2.5 预测。核心贡献可以表述为：在保持与 Transformer 相当或略好的预测精度时，Mamba 显著降低训练时间；同时分析不同聚合方法在站点异质性下的适用条件。"
    "如果后续多种子结果仍显示不到 0.2% 的精度差异，论文应强调效率与可扩展性，而不应把主要贡献写成精度提升。",
)

add_heading(doc, "建议形成的研究问题", 2)
add_numbered(doc, [
    "在多站点小时级 PM2.5 联邦预测中，Mamba 能否达到与 Transformer 相当的预测精度并降低训练成本。",
    "Mamba 的效率优势是否会随回看长度和预测范围增加而扩大。",
    "当监测站之间的数据分布差异增强时，FedProx 是否比 FedAvg 更稳定。",
    "不同监测站的样本量 缺失率 污染水平和周期性是否能解释模型收益差异。",
])

add_heading(doc, "可以直接检验的假设", 2)
add_table(
    doc,
    ["假设", "当前证据", "需要的验证"],
    [
        ["Mamba 在相近精度下显著缩短训练时间", "单轮快约4.6至4.8倍 总时长减少约77%至79%", "多种子 同硬件 参数量与精度对齐后的速度比较"],
        ["Mamba 加 FedAvg 是当前最佳组合", "最终 MAE RMSE 和 R²均为四组最佳", "站点级配对检验和独立验证集"],
        ["FedProx 的收益取决于客户端异质性", "当前设置下收益不稳定", "μ 本地轮数和异质性分层实验"],
        ["更长序列更有利于 Mamba", "当前168小时已经显示速度优势", "多个回看长度的误差 时间和显存曲线"],
    ],
    [2.45, 2.25, 2.25],
    font_size=8.8,
)

add_heading(doc, "论文创新性需要进一步加强的方向", 2)
add_body(
    doc,
    "仅将 Transformer 替换为 Mamba 并比较 FedAvg 与 FedProx，可能不足以形成强方法创新。"
    "更有潜力的方向是把空气质量的周期结构和站点差异纳入聚合机制，例如共享低频周期信息，同时保留站点特有的高频变化，或者对不同站点采用个性化预测头。"
    "这类设计能够把模型效率与空气质量问题的结构联系起来，比单纯增加模型数量更容易形成清晰贡献。",
)

add_heading(doc, "七 与教授沟通时的简明表述", 1)
add_body(
    doc,
    "目前我们完成了 Transformer 和 Mamba 分别结合 FedAvg 与 FedProx 的四组实验。"
    "在13个站点、回看168小时并预测24小时的设置下，四种方法的最终 MAE 都在 1.90 左右，精度差距很小。"
    "Mamba 加 FedAvg 的结果最好，最终 MAE 为 1.9021，曲线最低值为 1.8992。更明显的差别在训练效率：Mamba 每轮约35至39秒，Transformer 每轮约168至178秒，因此 Mamba 快约4.6至4.8倍。",
)
add_body(
    doc,
    "我认为现在可以把 Mamba 作为主要方向，但论文结论还需要进一步验证。"
    "当前结果看起来只有一个随机种子，而且训练过程中使用测试集画收敛曲线。下一步我准备先建立独立验证集，完成至少三个随机种子，并做站点级统计检验。"
    "之后再补充集中式 本地独立和朴素预测基线，以及不同序列长度和 FedProx 参数实验。",
)
add_body(
    doc,
    "我希望和您确认两点。第一，论文是否应以计算效率为主线，而不是强调很小的精度提升。第二，方法创新是否应继续发展为结合频率信息或站点个性化的聚合方法。",
)

add_heading(doc, "建议请教授决定的问题", 2)
add_numbered(doc, [
    "论文主线是否确定为在相近精度下提高联邦时序预测效率。",
    "是否需要在当前阶段加入新的聚合方法 还是先完成多种子和统计验证。",
    "论文目标期刊或会议对方法创新和实验规模的要求是什么。",
    "是否把频率感知聚合或个性化站点模型作为主要方法贡献。",
])

add_heading(doc, "八 阶段性结论", 1)
add_body(
    doc,
    "当前实验已经形成有价值的方向性证据。Mamba 加 FedAvg 在四组方案中具有最佳综合表现，尤其是训练时间优势非常明显。"
    "但现有精度差异不足以单独支撑强结论，且测试集使用方式 多种子复现 模型公平性和基线完整性仍需补充。"
    "下一阶段应优先完成验证集 多种子 站点级统计和复现审计。若这些结果继续成立，论文可以稳健地提出 Mamba 是多站点联邦 PM2.5 预测中更高效的骨干网络，并进一步研究针对空气质量周期性和站点差异的聚合方法。",
)

add_heading(doc, "附录 指标口径说明", 1)
add_table(
    doc,
    ["术语", "本报告中的含义"],
    [
        ["最终13站平均 MAE", "第30轮模型在13个站点测试集上的 MAE 算术平均"],
        ["曲线最低 MAE", "每5轮测试一次后 所有观测点中的最低13站平均测试 MAE"],
        ["进入最优值1%区间", "第一次达到不高于全程最低 MAE 乘以1.01的评估轮次"],
        ["每轮耗时", "客户端本地训练和参数聚合的平均时间 不完全包含评估与初始化"],
        ["总时长", "从构建客户端到完成最终评估的整体墙钟时间"],
    ],
    [2.05, 4.85],
    font_size=9.2,
)

# Keep table rows together where possible and ensure table borders are deterministic.
for table in doc.tables:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell in row.cells:
            set_cell_borders(cell)

doc.save(OUT_PATH)
print(OUT_PATH)
